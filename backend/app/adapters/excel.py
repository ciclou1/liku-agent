from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.adapters.generic import GenericAdapter
from app.exceptions import ValidationError


class ExcelAdapter(GenericAdapter):
    def __init__(self, required_columns: set[str]) -> None:
        super().__init__("excel")
        self.required_columns = required_columns

    def read_file(self, path: Path) -> list[dict[str, Any]]:
        if not path.exists() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValidationError("Excel 文件不存在或格式不支持")
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, OSError) as exc:
            raise ValidationError("Excel 文件内容无效") from exc
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValidationError("Excel 文件为空")
        headers = [str(value).strip() for value in rows[0] if value is not None]
        missing = self.required_columns - set(headers)
        if missing:
            raise ValidationError(f"Excel 缺少必填列: {', '.join(sorted(missing))}")
        return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]
